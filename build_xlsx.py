# -*- coding: utf-8 -*-
import pickle, datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils.dataframe import dataframe_to_rows

OUT = "/sessions/intelligent-festive-noether/mnt/outputs"

with open(f"{OUT}/eurotrip_data.pkl","rb") as f:
    data = pickle.load(f)

df_itinerario = data["itinerario"]
df_transportes = data["transportes"]
df_hospedagem = data["hospedagem"]
df_atracoes = data["atracoes"]
df_financeiro = data["financeiro"]
df_checklist = data["checklist"]

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1F3864", bold=True, size=16)
SUB_FONT = Font(color="595959", italic=True, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CONFIRM_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
SUGGEST_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

wb = Workbook()
wb.remove(wb.active)

def write_sheet(name, df, emoji_title, col_widths=None, date_cols=None, time_cols=None):
    ws = wb.create_sheet(name)
    ws.append([emoji_title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.append([f"EUROTRIP 2026 — {len(df)} registros"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(df.columns))
    ws.cell(row=2, column=1).font = SUB_FONT
    ws.append([])  # blank row
    header_row = 4
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=header_row, column=j, value=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for i, row in enumerate(df.itertuples(index=False), start=header_row+1):
        for j, val in enumerate(row, start=1):
            col_name = df.columns[j-1]
            cell = ws.cell(row=i, column=j)
            if date_cols and col_name in date_cols and val:
                try:
                    cell.value = datetime.datetime.strptime(str(val), "%Y-%m-%d").date()
                    cell.number_format = "DD/MM/YYYY"
                except Exception:
                    cell.value = val
            elif time_cols and col_name in time_cols and val:
                try:
                    cell.value = datetime.datetime.strptime(str(val), "%H:%M").time()
                    cell.number_format = "HH:MM"
                except Exception:
                    cell.value = val
            else:
                cell.value = val
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col_name in ("Descrição","Observações")))
    ws.freeze_panes = ws.cell(row=header_row+1, column=1)
    widths = col_widths or [18]*len(df.columns)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    last_row = header_row + len(df)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(df.columns))}{last_row}"
    return ws, header_row, last_row

# ---------------- 1) ITINERARIO ----------------
ws1, hr1, lr1 = write_sheet("1_Itinerario", df_itinerario, "🗺️ ITINERÁRIO",
    col_widths=[8,12,16,14,34,15,45,18], date_cols=["Data"])
status_col = df_itinerario.columns.get_loc("Status") + 1
dv1 = DataValidation(type="list", formula1='"Confirmado,Sugestão - validar,Concluído,Cancelado"', allow_blank=True)
ws1.add_data_validation(dv1)
dv1.add(f"{get_column_letter(status_col)}{hr1+1}:{get_column_letter(status_col)}{lr1}")
col_l = get_column_letter(status_col)
ws1.conditional_formatting.add(f"{col_l}{hr1+1}:{col_l}{lr1}",
    CellIsRule(operator="equal", formula=['"Confirmado"'], fill=CONFIRM_FILL))
ws1.conditional_formatting.add(f"{col_l}{hr1+1}:{col_l}{lr1}",
    CellIsRule(operator="equal", formula=['"Sugestão - validar"'], fill=SUGGEST_FILL))
tipo_col = df_itinerario.columns.get_loc("Tipo") + 1
dv1b = DataValidation(type="list", formula1='"Turismo,Deslocamento,Descanso"', allow_blank=True)
ws1.add_data_validation(dv1b)
dv1b.add(f"{get_column_letter(tipo_col)}{hr1+1}:{get_column_letter(tipo_col)}{lr1}")

# ---------------- 2) TRANSPORTES ----------------
ws2, hr2, lr2 = write_sheet("2_Transportes", df_transportes, "🚗 TRANSPORTES",
    col_widths=[8,10,16,20,12,11,11,20,10], date_cols=["Data"], time_cols=["Hora Saída","Hora Chegada"])
# formula real de duração (sobrescreve a coluna Duração com fórmula Excel)
dur_col = df_transportes.columns.get_loc("Duração") + 1
sai_col = df_transportes.columns.get_loc("Hora Saída") + 1
che_col = df_transportes.columns.get_loc("Hora Chegada") + 1
for i in range(hr2+1, lr2+1):
    sai = get_column_letter(sai_col) + str(i)
    che = get_column_letter(che_col) + str(i)
    formula = f'=IF(OR({sai}="",{che}=""),"",TEXT(MOD({che}-{sai},1),"[h]\\"h\\"mm"))'
    ws2.cell(row=i, column=dur_col).value = formula
tipo_col2 = df_transportes.columns.get_loc("Tipo") + 1
dv2 = DataValidation(type="list", formula1='"Trem,Voo,Carro,Barco,Ônibus"', allow_blank=True)
ws2.add_data_validation(dv2)
dv2.add(f"{get_column_letter(tipo_col2)}{hr2+1}:{get_column_letter(tipo_col2)}{lr2}")
status_col2 = df_transportes.columns.get_loc("Status") + 1
dv2b = DataValidation(type="list", formula1='"Confirmado,Sugestão - validar,Comprado,Cancelado"', allow_blank=True)
ws2.add_data_validation(dv2b)
dv2b.add(f"{get_column_letter(status_col2)}{hr2+1}:{get_column_letter(status_col2)}{lr2}")

# ---------------- 3) HOSPEDAGEM ----------------
ws3, hr3, lr3 = write_sheet("3_Hospedagem", df_hospedagem, "🏨 HOSPEDAGEM",
    col_widths=[8,16,32,12,12,14,45], date_cols=["Check-in","Check-out"])
noites_col = len(df_hospedagem.columns) + 1
ws3.cell(row=hr3, column=noites_col, value="Noites").fill = HEADER_FILL
ws3.cell(row=hr3, column=noites_col).font = HEADER_FONT
ws3.cell(row=hr3, column=noites_col).border = BORDER
ci_col = df_hospedagem.columns.get_loc("Check-in") + 1
co_col = df_hospedagem.columns.get_loc("Check-out") + 1
for i in range(hr3+1, lr3+1):
    ci = get_column_letter(ci_col) + str(i)
    co = get_column_letter(co_col) + str(i)
    ws3.cell(row=i, column=noites_col).value = f"={co}-{ci}"
    ws3.cell(row=i, column=noites_col).border = BORDER
ws3.column_dimensions[get_column_letter(noites_col)].width = 10

# ---------------- 4) ATRACOES ----------------
ws4, hr4, lr4 = write_sheet("4_Atracoes", df_atracoes, "🎟️ ATRAÇÕES",
    col_widths=[8,32,16,16,16,16,26,14])
ing_col = df_atracoes.columns.get_loc("Necessita Ingresso") + 1
dv4 = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ws4.add_data_validation(dv4)
dv4.add(f"{get_column_letter(ing_col)}{hr4+1}:{get_column_letter(ing_col)}{lr4}")
tipo_col4 = df_atracoes.columns.get_loc("Tipo") + 1
dv4b = DataValidation(type="list", formula1='"Museu,Ponto turístico,Passeio,Memorial/Museu"', allow_blank=True)
ws4.add_data_validation(dv4b)
dv4b.add(f"{get_column_letter(tipo_col4)}{hr4+1}:{get_column_letter(tipo_col4)}{lr4}")

# ---------------- 5) FINANCEIRO ----------------
ws5, hr5, lr5 = write_sheet("5_Financeiro", df_financeiro, "💸 CONTROLE FINANCEIRO",
    col_widths=[8,16,46,14,14,10,10])
ve_col = df_financeiro.columns.get_loc("Valor Estimado") + 1
for i in range(hr5+1, lr5+1):
    ws5.cell(row=i, column=ve_col).number_format = "#,##0.00"
pago_col = df_financeiro.columns.get_loc("Pago?") + 1
dv5 = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ws5.add_data_validation(dv5)
dv5.add(f"{get_column_letter(pago_col)}{hr5+1}:{get_column_letter(pago_col)}{lr5}")
cat_col = df_financeiro.columns.get_loc("Categoria") + 1
dv5b = DataValidation(type="list", formula1='"Transporte,Alimentação,Hospedagem,Passeio"', allow_blank=True)
ws5.add_data_validation(dv5b)
dv5b.add(f"{get_column_letter(cat_col)}{hr5+1}:{get_column_letter(cat_col)}{lr5}")
# linha de total
tot_row = lr5 + 2
ws5.cell(row=tot_row, column=cat_col-1, value="TOTAL ESTIMADO (EUR)").font = Font(bold=True)
ws5.cell(row=tot_row, column=ve_col).value = f"=SUM({get_column_letter(ve_col)}{hr5+1}:{get_column_letter(ve_col)}{lr5})"
ws5.cell(row=tot_row, column=ve_col).font = Font(bold=True)
ws5.cell(row=tot_row, column=ve_col).number_format = "#,##0.00"

# ---------------- 6) CHECKLIST ----------------
ws6, hr6, lr6 = write_sheet("6_Checklist", df_checklist, "🎒 CHECKLIST DE VIAGEM", col_widths=[18,55,14])
status_col6 = df_checklist.columns.get_loc("Status") + 1
dv6 = DataValidation(type="list", formula1='"Pendente,OK,Comprado"', allow_blank=True)
ws6.add_data_validation(dv6)
dv6.add(f"{get_column_letter(status_col6)}{hr6+1}:{get_column_letter(status_col6)}{lr6}")
ws6.conditional_formatting.add(f"{get_column_letter(status_col6)}{hr6+1}:{get_column_letter(status_col6)}{lr6}",
    CellIsRule(operator="equal", formula=['"OK"'], fill=CONFIRM_FILL))
ws6.conditional_formatting.add(f"{get_column_letter(status_col6)}{hr6+1}:{get_column_letter(status_col6)}{lr6}",
    CellIsRule(operator="equal", formula=['"Pendente"'], fill=SUGGEST_FILL))

# ---------------- 0) RESUMO (capa) ----------------
ws0 = wb.create_sheet("0_Resumo", 0)
ws0.merge_cells("A1:D1")
ws0["A1"] = "EUROTRIP 2026 — Painel de Controle"
ws0["A1"].font = Font(bold=True, size=20, color="1F3864")
ws0.merge_cells("A2:D2")
ws0["A2"] = "Itália · França · Luxemburgo · Alemanha · Polônia · Albânia — 30 dias (06/08 a 04/09/2026)"
ws0["A2"].font = SUB_FONT

resumo = [
    ("Total de dias de viagem", 30),
    ("Países visitados", 6),
    ("Atividades no itinerário", len(df_itinerario)),
    ("Trechos de transporte", len(df_transportes)),
    ("Hospedagens planejadas", len(df_hospedagem)),
    ("Atrações mapeadas", len(df_atracoes)),
    ("Itens no checklist", len(df_checklist)),
    ("Dias confirmados (documento original)", (df_itinerario["Status"]=="Confirmado").sum()),
    ("Dias em rascunho (sugestão a validar)", (df_itinerario["Status"]=="Sugestão - validar").sum()),
]
r = 4
ws0.cell(row=r, column=1, value="Métrica").font = HEADER_FONT
ws0.cell(row=r, column=1).fill = HEADER_FILL
ws0.cell(row=r, column=2, value="Valor").font = HEADER_FONT
ws0.cell(row=r, column=2).fill = HEADER_FILL
for label, val in resumo:
    r += 1
    ws0.cell(row=r, column=1, value=label).border = BORDER
    ws0.cell(row=r, column=2, value=val).border = BORDER
r += 2
ws0.cell(row=r, column=1, value="Orçamento total estimado (EUR)").font = Font(bold=True)
ws0.cell(row=r, column=2).value = f"=SUM('5_Financeiro'!D5:D{hr5+len(df_financeiro)})"
ws0.cell(row=r, column=2).number_format = "#,##0.00"
ws0.cell(row=r, column=2).font = Font(bold=True, size=12, color="1F3864")
r += 2
ws0.cell(row=r, column=1, value="Legenda de status:").font = Font(bold=True)
r += 1
ws0.cell(row=r, column=1, value="Confirmado / OK").fill = CONFIRM_FILL
r += 1
ws0.cell(row=r, column=1, value="Sugestão - validar / Pendente").fill = SUGGEST_FILL
ws0.column_dimensions["A"].width = 42
ws0.column_dimensions["B"].width = 18

wb.save(f"{OUT}/EUROTRIP_2026_Sistema.xlsx")
print("xlsx salvo com sucesso")
